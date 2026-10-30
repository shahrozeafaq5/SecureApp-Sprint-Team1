import os
import sys
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client

ROOT = Path(__file__).resolve().parent
SCREENSHOTS_DIR = ROOT / "screenshots"

SCREENSHOT_MAP = {
    "16": "screenshot_16_buffer_overflow_before.png",
    "17": "screenshot_17_buffer_overflow_after.png",
    "18": "screenshot_18_command_injection_before.png",
    "19": "screenshot_19_command_injection_after.png",
    "26": "screenshot_26_automated_tests_all_pass.png",
    "27": "screenshot_27_bandit_sast_before.png",
    "28": "screenshot_28_bandit_sast_after.png",
    "29": "screenshot_29_cicd_pipeline_success.png",
    "30": "screenshot_30_jira_board.png",
}

def markdown_to_docx(md_path: Path, docx_path: Path, title: str):
    doc = docx.Document()
    
    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("COMSATS University Islamabad | CYC386 Secure Software Design and Development\nTeam: Talal Nadeem Awan (FA23-BCT-051) & Shahroze Afaq (FA23-BCT-039)")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_code_block = False
    code_text = []
    
    for line in lines:
        stripped = line.rstrip()
        
        # Code block toggle
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run("\n".join(code_text))
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(15, 23, 42)
                code_text = []
            else:
                in_code_block = True
                code_text = []
            continue
            
        if in_code_block:
            code_text.append(stripped)
            continue
            
        if not stripped:
            continue
            
        # Headers
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:])
            r.font.name = 'Arial'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[3:])
            r.font.name = 'Arial'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(stripped[4:])
            r.font.name = 'Arial'
            r.font.size = Pt(11)
            r.font.bold = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(stripped[2:])
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(2)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run(stripped[2:])
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.color.rgb = RGBColor(100, 116, 139)
            
            # Check for screenshot embedding
            match = re.search(r"SCREENSHOT CHECKPOINT #(\d+)", stripped)
            if match:
                cp_num = match.group(1)
                img_name = SCREENSHOT_MAP.get(cp_num)
                if img_name:
                    img_path = SCREENSHOTS_DIR / img_name
                    if img_path.is_file():
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(8)
                        run_img = p_img.add_run()
                        run_img.add_picture(str(img_path), width=Inches(5.8))
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run(f"Figure: Authentic Evidence for Checkpoint #{cp_num}")
                        r_cap.font.name = 'Arial'
                        r_cap.font.size = Pt(9)
                        r_cap.font.italic = True
        else:
            p = doc.add_paragraph()
            r = p.add_run(stripped)
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            
    doc.save(str(docx_path))
    print(f"[+] Saved Word document with embedded screenshots: {docx_path}")

def convert_docx_to_pdf(docx_path: Path, pdf_path: Path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        doc.Close()
        print(f"[+] Converted to PDF: {pdf_path}")
    finally:
        word.Quit()

def main():
    tasks = [
        ("SECURITY_IMPLEMENTATION.md", "SECURITY_IMPLEMENTATION.docx", "SECURITY_IMPLEMENTATION.pdf", "Security Implementation & Vulnerability Remediation"),
        ("FINAL_REPORT_IEEE.md", "FINAL_REPORT_IEEE.docx", "FINAL_REPORT_IEEE.pdf", "Attack-Resistant Document Vault: DevSecOps Engineering Sprint"),
        ("PNE_Report.md", "PNE_Report.docx", "PNE_Report.pdf", "Protection Needs Elicitation (PNE) & CIA Mapping"),
        ("docs/THREAT_MODEL.md", "docs/THREAT_MODEL.docx", "docs/THREAT_MODEL.pdf", "Threat Modeling & Quantitative Risk Assessment"),
    ]
    
    for md_rel, docx_rel, pdf_rel, title in tasks:
        md_file = ROOT / md_rel
        docx_file = ROOT / docx_rel
        pdf_file = ROOT / pdf_rel
        
        if md_file.is_file():
            print(f"[*] Processing {md_rel}...")
            markdown_to_docx(md_file, docx_file, title)
            convert_docx_to_pdf(docx_file, pdf_file)

if __name__ == "__main__":
    main()
